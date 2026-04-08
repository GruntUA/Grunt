"""Export and Print operations for Documents."""

from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from typing import TYPE_CHECKING, Any

import openpyxl
import structlog
from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import Response, StreamingResponse
from openpyxl.styles import Font, PatternFill

from grunt.api.v1.docs.utils import (
    _fmt,
    _generate_xlsx_single,
    _non_layout_fields,
    get_doc_service,
)
from grunt.core.auth.dependencies import current_user
from grunt.core.db.session import get_session
from grunt.core.metadata.registry import doctype_registry

if TYPE_CHECKING:
    from sqlalchemy.ext.asyncio import AsyncSession

    from grunt.core.auth.models import GruntUser
    from grunt.core.document.service import DocumentService

logger = structlog.get_logger()
router = APIRouter()


@router.get("/{doctype}/export")
async def export_documents(
    doctype: str,
    user: GruntUser = Depends(current_user),
    svc: DocumentService = Depends(get_doc_service),
    sort_by: str = "modified_at",
    sort_order: str = "desc",
    filters: str | None = None,
    search: str | None = None,
    fields: str | None = None,
) -> StreamingResponse:
    """Export documents as CSV."""
    parsed_filters: dict[str, Any] = {}
    if filters:
        try:
            parsed_filters = json.loads(filters)
        except json.JSONDecodeError:
            logger.debug("export_invalid_filters", doctype=doctype)

    parsed_fields = fields.split(",") if fields else None

    res = await svc.list_documents(
        doctype,
        user,
        per_page=1000,
        sort_by=sort_by,
        sort_order=sort_order,
        filters=parsed_filters,
        search=search,
        fields=parsed_fields,
    )
    data = res["data"]

    # 2. Generate CSV
    output = io.StringIO()
    if data:
        writer = csv.DictWriter(output, fieldnames=data[0].keys())
        writer.writeheader()
        writer.writerows(data)

    output.seek(0)

    filename = f"{doctype}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"

    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/{doctype}/{doc_id}/print")
async def print_document(
    doctype: str,
    doc_id: str,
    fmt: str = Query("html"),  # html | pdf | xlsx | docx
    print_format: str | None = Query(None),
    user: GruntUser = Depends(current_user),
    svc: DocumentService = Depends(get_doc_service),
    session: AsyncSession = Depends(get_session),
) -> Response:
    """Generate document in requested format: html, pdf, xlsx, or docx."""
    doc = await svc.get_document(doctype, doc_id, user)
    dt = await doctype_registry.get(doctype)

    if fmt == "xlsx":
        content = _generate_xlsx_single(dt, doc)
        return Response(
            content=content,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{doctype}_{doc_id[:8]}.xlsx"'},
        )

    if fmt == "docx":
        from grunt.core.print.renderer import get_print_format_template  # noqa: PLC0415

        pf = await get_print_format_template(session, doctype, print_format)
        if pf and pf[1] == "docx":
            from grunt.core.print.renderer import render_docx  # noqa: PLC0415

            try:
                docx_bytes = render_docx(pf[0], doc)
                return Response(
                    content=docx_bytes,
                    media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={
                        "Content-Disposition": f'attachment; filename="{doctype}_{doc_id[:8]}.docx"'
                    },
                )
            except Exception as e:
                raise HTTPException(status_code=500, detail=f"DOCX generation error: {e}")
        raise HTTPException(
            status_code=404,
            detail="Шаблон DOCX не знайдено. Створіть PrintFormat з template_type='docx'.",
        )

    if fmt in ("pdf", "html"):
        from grunt.core.print.renderer import (  # noqa: PLC0415
            get_print_format_template,
            render_from_string,
            render_standard,
        )

        html = None

        # 1. Try custom PrintFormat from DB
        pf = await get_print_format_template(session, doctype, print_format)
        if pf and pf[1] == "html":
            try:
                html = render_from_string(pf[0], doc, doctype_label=dt.label, fields=dt.fields)
            except Exception:
                pass  # Fall back to standard

        # 2. Standard template (auto-generated from fields)
        if html is None:
            html = render_standard(dt.label, dt.fields, doc)

        if fmt == "html":
            return Response(content=html, media_type="text/html")

        # PDF via WeasyPrint
        try:
            from weasyprint import HTML as WeasyprintHTML  # noqa: PLC0415

            pdf_bytes = WeasyprintHTML(string=html).write_pdf()
            return Response(
                content=pdf_bytes,
                media_type="application/pdf",
                headers={
                    "Content-Disposition": f'attachment; filename="{doctype}_{doc_id[:8]}.pdf"'
                },
            )
        except ImportError:
            raise HTTPException(
                status_code=501,
                detail="WeasyPrint не встановлено. Використайте fmt=xlsx або fmt=html",
            )

    raise HTTPException(
        status_code=400,
        detail=f"Невідомий формат: {fmt}. Підтримуються: html, pdf, xlsx, docx",
    )


@router.post("/{doctype}/print-preview")
async def print_format_preview(
    doctype: str,
    body: dict[str, Any],
    user: GruntUser = Depends(current_user),
    svc: DocumentService = Depends(get_doc_service),
) -> Response:
    """Render an arbitrary Jinja2 HTML template against a document (or sample)."""
    from grunt.core.print.renderer import render_from_string  # noqa: PLC0415

    template_str: str = (body.get("template") or "").strip()
    if not template_str:
        raise HTTPException(422, detail="template is required")

    doc_id: str | None = body.get("doc_id")
    dt = await doctype_registry.get(doctype)

    if doc_id:
        doc = await svc.get_document(doctype, doc_id, user)
    else:
        # Use first available document as sample
        result = await svc.list_documents(doctype, user=user, page=1, per_page=1)
        sample_list = result.get("data") or []
        if sample_list:
            doc = dict(sample_list[0])
        else:
            # Synthetic empty doc with all field names set to None
            doc = {f.fieldname: None for f in dt.fields}
            doc.setdefault("id", "preview")
            doc.setdefault("name", "Зразок")

    try:
        html = render_from_string(template_str, doc, doctype_label=dt.label, fields=dt.fields)
    except Exception as e:
        html = f"<pre style='color:red;padding:1rem'>Помилка шаблону:\n{e}</pre>"

    return Response(content=html, media_type="text/html")


@router.get("/{doctype}/export/xlsx")
async def export_documents_xlsx(
    doctype: str,
    request: Request,
    user: GruntUser = Depends(current_user),
    svc: DocumentService = Depends(get_doc_service),
) -> Response:
    """Export all documents (up to 10 000) for a DocType as XLSX."""
    filters: dict[str, str] = {}
    for key, value in request.query_params.items():
        if key.startswith("filter[") and key.endswith("]"):
            filters[key[7:-1]] = value

    result = await svc.list_documents(
        doctype_name=doctype,
        user=user,
        page=1,
        per_page=10000,
        filters=filters if filters else None,
    )
    rows: list[dict[str, Any]] = result["data"]
    dt = await doctype_registry.get(doctype)
    visible = _non_layout_fields(dt)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = dt.label[:31]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2D6A4F")

    # Header row
    headers = ["ID", "Назва", "Власник", "Створено"] + [f.label for f in visible]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill

    # Data rows
    for row_idx, row in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=str(row.get("id", "")))
        ws.cell(row=row_idx, column=2, value=str(row.get("name", "")))
        ws.cell(row=row_idx, column=3, value=str(row.get("owner", "")))
        ws.cell(row=row_idx, column=4, value=_fmt(row.get("created_at")))
        for col_idx, field in enumerate(visible, start=5):
            ws.cell(row=row_idx, column=col_idx, value=_fmt(row.get(field.fieldname)))

    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{doctype}_export.xlsx"'},
    )
