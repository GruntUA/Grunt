"""Utility functions and dependencies for Document API."""

from __future__ import annotations

import io
from datetime import date, datetime
from typing import TYPE_CHECKING, Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from grunt.document.meta import Meta
from grunt.io.exporters.sanitize import escape_formula

if TYPE_CHECKING:
    from fastapi import Request



def parse_query_filters(request: Request) -> dict[str, str]:
    """Merge ``quick_filter[field__op]=value`` and ``filter[field__op]=value`` query params.

    ``quick_filter[...]`` is applied first (lower precedence); an explicit
    ``filter[...]`` for the same key overrides it. Shared by the three
    endpoints that accept both styles from the query string: document list,
    tree, and tree children.
    """
    merged: dict[str, str] = {}
    qf_prefix = "quick_filter["
    for key, value in request.query_params.items():
        if key.startswith(qf_prefix) and key.endswith("]"):
            merged[key[len(qf_prefix) : -1]] = value
    for key, value in request.query_params.items():
        if key.startswith("filter[") and key.endswith("]"):
            merged[key[7:-1]] = value
    return merged


def _fmt(val: object) -> str:
    """Format a value for human display."""
    if val is None:
        return ""
    if isinstance(val, bool):
        return "Так" if val else "Ні"
    if isinstance(val, datetime):
        return val.strftime("%d.%m.%Y %H:%M")
    if isinstance(val, date):
        return val.strftime("%d.%m.%Y")
    return escape_formula(str(val))


def _non_layout_fields(dt: Any) -> list[Any]:
    """Return fields that are not layout-only (Section, Column, Tab)."""
    return Meta(dt).get_visible_fields()


def _generate_xlsx_single(dt: Any, doc: dict[str, Any]) -> bytes:
    """Generate a single document XLSX."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = dt.label[:31]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2D6A4F")

    ws.merge_cells("A1:B1")
    title_cell = ws["A1"]
    title_cell.value = dt.label
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    ws["A2"] = "Поле"
    ws["B2"] = "Значення"
    for cell in [ws["A2"], ws["B2"]]:
        cell.font = header_font
        cell.fill = header_fill

    for i, field in enumerate(_non_layout_fields(dt), start=3):
        ws[f"A{i}"] = field.label
        ws[f"B{i}"] = _fmt(doc.get(field.fieldname))

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 50

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _generate_html_single(dt: Any, doc: dict[str, Any]) -> str:
    """Generate a single document HTML (standard layout).

    Hand-built (not Jinja, which autoescapes) — every interpolated value is
    document data or a label and MUST be escaped explicitly, same reasoning
    as print/renderer.py's _render_fallback (a field value containing HTML
    is document data, not markup, and this doesn't get autoescape for free).
    """
    from html import escape

    rows = ""
    for field in _non_layout_fields(dt):
        val = _fmt(doc.get(field.fieldname))
        rows += f"<tr><th>{escape(field.label)}</th><td>{escape(val)}</td></tr>\n"

    created = escape(_fmt(doc.get("created_at")))
    owner = escape(str(doc.get("owner", "")))

    return f"""<!DOCTYPE html>
<html>
<head>
<meta charset="UTF-8">
<style>
  body {{ font-family: Arial, sans-serif; font-size: 12pt; margin: 2cm; }}
  h1 {{ color: #2D6A4F; border-bottom: 2px solid #2D6A4F; padding-bottom: 8px; }}
  table {{ width: 100%; border-collapse: collapse; margin-top: 16px; }}
  th {{ text-align: left; padding: 8px 12px; background: #f0f4f0; font-weight: bold; width: 40%; }}
  td {{ padding: 8px 12px; border-bottom: 1px solid #e0e0e0; }}
  .meta {{ color: #666; font-size: 10pt; margin-top: 24px; }}
</style>
</head>
<body>
  <h1>{escape(dt.label)}</h1>
  <table>{rows}</table>
  <div class="meta">Створено: {created} | Автор: {owner}</div>
</body>
</html>"""
