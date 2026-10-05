"""XLSX exporter - produces a styled Excel workbook."""

from __future__ import annotations

import io
from datetime import date, datetime
from typing import TYPE_CHECKING, Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from grunt.io.exporters.registry import Exporter
from grunt.io.exporters.sanitize import escape_formula

if TYPE_CHECKING:
    from grunt.metadata.field import DocField


# Header style
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="2D6A4F")
_HEADER_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=False)


def _fmt(value: Any) -> Any:
    """Coerce a value into something openpyxl can write natively."""
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, dict):
        # e.g. Geolocation {"lat": ..., "lng": ...}
        if "lat" in value and "lng" in value:
            return escape_formula(f"{value['lat']}, {value['lng']}")
        return escape_formula(str(value))
    return escape_formula(str(value))


class XlsxExporter(Exporter):
    id = "xlsx"
    label = "Excel (XLSX)"
    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    file_extension = "xlsx"

    async def export(
        self,
        doctype: str,
        rows: list[dict[str, Any]],
        fields: list[DocField],
    ) -> bytes:
        exportable = [f for f in fields if f.is_physical]

        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None  # a new Workbook always has one sheet
        ws.title = doctype[:31]  # Excel sheet name limit

        # Header row
        headers = [f.label for f in exportable]
        for col_idx, label in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = _HEADER_ALIGN

        # Data rows
        for row_idx, row in enumerate(rows, start=2):
            for col_idx, field in enumerate(exportable, start=1):
                ws.cell(
                    row=row_idx,
                    column=col_idx,
                    value=_fmt(row.get(field.fieldname)),
                )

        # Auto-fit column widths (approximate)
        for col_idx, col_cells in enumerate(ws.columns, start=1):
            max_len = max((len(str(c.value or "")) for c in col_cells), default=0)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
