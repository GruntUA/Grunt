"""Tests for the Reports module.

Plain CRUD (list/get/create/update/delete) on the ``Report`` doctype is
covered end-to-end through the generic ``/api/v1/docs/Report`` REST routes —
these tests exercise real permission enforcement (superadmin-only
write/create/delete, any-authenticated-user read) rather than calling
Python functions directly under a superadmin-bypassing context.

Report *execution* (``run``/``preview``) and the xlsx export
(``export_xlsx``) are RPC-only — those are still exercised as direct
function calls under the ``ctx`` fixture (SYSTEM_USER), matching the
pre-existing style for that part of the module.
"""

from __future__ import annotations

import pytest
from httpx import AsyncClient


async def _register_regular_user(client: AsyncClient, email: str) -> dict[str, str]:
    """Register a second (non-superadmin) user and return its auth headers.

    The first user ever registered on a site becomes superadmin; any
    subsequent registration is a plain user (see grunt/auth/doctypes/User/user.py).
    """
    r_reg = await client.post(
        "/api/v1/method/grunt.auth.doctypes.User.user.register_full_name_api",
        json={"email": email, "password": "secret", "full_name": "Regular User"},
    )
    assert r_reg.status_code in (200, 201, 409)

    resp = await client.post(
        "/api/v1/method/grunt.auth.doctypes.User.user.login_api",
        json={"email": email, "password": "secret"},
    )
    assert resp.status_code == 200, f"Auth failed: {resp.text}"
    token = resp.json()["data"]["access_token"]
    return {"Authorization": f"Bearer {token}"}


# ── Generic docs CRUD: permissions ──────────────────────────────────────


@pytest.mark.asyncio
async def test_report_crud_full_cycle_as_superadmin(client: AsyncClient, auth_headers):
    """Superadmin can list/create/get/update/delete a Report via generic docs CRUD."""
    # Create
    r_create = await client.post(
        "/api/v1/docs/Report",
        json={
            "report_name": "Users Report",
            "report_type": "Query",
            "query": "SELECT 1 as num",
        },
        headers=auth_headers,
    )
    assert r_create.status_code == 201, r_create.text
    report_id = r_create.json()["data"]["name"]

    # List
    r_list = await client.get("/api/v1/docs/Report", headers=auth_headers)
    assert r_list.status_code == 200
    assert any(r["report_name"] == "Users Report" for r in r_list.json()["data"])

    # Get one
    r_get = await client.get(f"/api/v1/docs/Report/{report_id}", headers=auth_headers)
    assert r_get.status_code == 200
    assert r_get.json()["data"]["report_name"] == "Users Report"

    # Update
    r_update = await client.put(
        f"/api/v1/docs/Report/{report_id}",
        json={"report_type": "Script", "script": "result = {'columns': [], 'data': []}"},
        headers=auth_headers,
    )
    assert r_update.status_code == 200, r_update.text
    assert r_update.json()["data"]["report_type"] == "Script"

    # Delete
    r_delete = await client.delete(f"/api/v1/docs/Report/{report_id}", headers=auth_headers)
    assert r_delete.status_code == 204

    r_get_after = await client.get(f"/api/v1/docs/Report/{report_id}", headers=auth_headers)
    assert r_get_after.status_code == 404


@pytest.mark.asyncio
async def test_report_read_allowed_for_any_authenticated_user(client: AsyncClient, auth_headers):
    """Any authenticated user can read reports, not just superadmin/System Manager."""
    r_create = await client.post(
        "/api/v1/docs/Report",
        json={"report_name": "Readable Report", "report_type": "Query", "query": "SELECT 1"},
        headers=auth_headers,
    )
    assert r_create.status_code == 201, r_create.text

    regular_headers = await _register_regular_user(client, "regular-reader@grunt.example.com")

    r_list = await client.get("/api/v1/docs/Report", headers=regular_headers)
    assert r_list.status_code == 200
    assert any(r["report_name"] == "Readable Report" for r in r_list.json()["data"])


@pytest.mark.asyncio
async def test_report_write_forbidden_for_regular_user(client: AsyncClient, auth_headers):
    """A regular (non-superadmin) user cannot create, update, or delete reports.

    Regression guard for the Report.json permissions weakening this test suite
    was written to catch: the doctype's permissions must not grant write to
    any role broader than superadmin (see Report.json — only {"role": "All",
    "read": true} is defined, so write/create/delete fall through to
    "nobody but superadmin").
    """
    # Seed one report as superadmin so there's something to attack via update/delete.
    r_create = await client.post(
        "/api/v1/docs/Report",
        json={"report_name": "Protected Report", "report_type": "Query", "query": "SELECT 1"},
        headers=auth_headers,
    )
    assert r_create.status_code == 201, r_create.text
    report_id = r_create.json()["data"]["name"]

    regular_headers = await _register_regular_user(client, "regular-writer@grunt.example.com")

    r_create_denied = await client.post(
        "/api/v1/docs/Report",
        json={"report_name": "Hacked Report", "report_type": "Query", "query": "SELECT 1"},
        headers=regular_headers,
    )
    assert r_create_denied.status_code == 403

    r_update_denied = await client.put(
        f"/api/v1/docs/Report/{report_id}",
        json={"report_type": "Script"},
        headers=regular_headers,
    )
    assert r_update_denied.status_code == 403

    r_delete_denied = await client.delete(
        f"/api/v1/docs/Report/{report_id}", headers=regular_headers
    )
    assert r_delete_denied.status_code == 403


# ── run_report / run_preview (RPC-only, unaffected by the CRUD migration) ──


@pytest.mark.asyncio
async def test_run_query_report(ctx):
    """Run a SELECT query report."""
    from grunt.reports.doctypes.Report.report import run as run_report

    await ctx.new_doc(
        "Report",
        {
            "report_name": "Select Report",
            "report_type": "Query",
            "query": "SELECT 42 as answer",
        },
    )
    await ctx.db._session().commit()

    data = await run_report(name="Select Report", filters={})
    assert data["data"][0]["answer"] == 42


@pytest.mark.asyncio
async def test_run_script_report_result_contract(ctx):
    """A Script report may build rows itself via `result = {...}`."""
    from grunt.reports.doctypes.Report.report import run as run_report

    await ctx.new_doc(
        "Report",
        {
            "report_name": "Script Result Report",
            "report_type": "Script",
            "script": (
                "result = {"
                "  'columns': [{'fieldname': 'n', 'label': 'N', 'fieldtype': 'Int'}],"
                "  'data': [{'n': 1}, {'n': 2}],"
                "}"
            ),
        },
    )
    await ctx.db._session().commit()

    data = await run_report(name="Script Result Report", filters={})
    assert [r["n"] for r in data["data"]] == [1, 2]


@pytest.mark.asyncio
async def test_run_script_report_query_contract(ctx):
    """A Script report may just pick a read-only SQL string per `db_dialect`."""
    from grunt.reports.doctypes.Report.report import run as run_report

    await ctx.new_doc(
        "Report",
        {
            "report_name": "Script Query Report",
            "report_type": "Script",
            "script": (
                "if db_dialect == 'postgresql':\n"
                "    query = 'SELECT 1 AS answer'\n"
                "else:\n"
                "    query = 'SELECT 42 AS answer'\n"
            ),
        },
    )
    await ctx.db._session().commit()

    data = await run_report(name="Script Query Report", filters={})
    assert data["data"][0]["answer"] == 42


@pytest.mark.asyncio
async def test_run_script_report_query_contract_still_select_only(ctx):
    """The `query =` escape hatch runs through the same SELECT-only guard."""
    from fastapi import HTTPException

    from grunt.api.messages import ApplicationError
    from grunt.errors import GruntError
    from grunt.reports.doctypes.Report.report import run as run_report

    await ctx.new_doc(
        "Report",
        {
            "report_name": "Script Query Danger Report",
            "report_type": "Script",
            "script": "query = 'DELETE FROM grunt_auth_user'",
        },
    )
    await ctx.db._session().commit()

    with pytest.raises((HTTPException, ApplicationError, GruntError)):
        await run_report(name="Script Query Danger Report", filters={})


@pytest.mark.asyncio
async def test_run_report_forbids_delete(ctx):
    """DELETE SQL is blocked in query reports."""
    from fastapi import HTTPException

    from grunt.reports.doctypes.Report.report import run as run_report

    await ctx.new_doc(
        "Report",
        {
            "report_name": "Bad Report",
            "report_type": "Query",
            "query": "DELETE FROM grunt_auth_user",
        },
    )
    await ctx.db._session().commit()

    from grunt.api.messages import ApplicationError
    from grunt.errors import GruntError

    with pytest.raises((HTTPException, ApplicationError, GruntError)):
        await run_report(name="Bad Report", filters={})


@pytest.mark.asyncio
async def test_run_list_report_operator_filters(ctx):
    """A List report applies operator-aware filters (``field__eq`` / ``__like`` / …)."""
    from grunt.reports.doctypes.Report.report import run as run_report

    for rn, rt in [("L Alpha", "Query"), ("L Beta", "Query"), ("L Gamma", "Script")]:
        await ctx.new_doc(
            "Report",
            {
                "report_name": rn,
                "report_type": rt,
                "query": "SELECT 1",
                "script": "result = {'data': []}",
            },
        )
    await ctx.new_doc(
        "Report",
        {
            "report_name": "List Over Reports",
            "report_type": "List",
            "doctype": "Report",
            "columns": [
                {"fieldname": "report_name", "label": "Name"},
                {"fieldname": "report_type", "label": "Type"},
            ],
            "filters_config": [
                {"fieldname": "report_type", "label": "Type", "fieldtype": "Select"},
                {"fieldname": "report_name", "label": "Name", "fieldtype": "Data"},
            ],
        },
    )
    await ctx.db._session().commit()

    eq_res = await run_report(name="List Over Reports", filters={"report_type__eq": "Script"})
    assert {r["report_name"] for r in eq_res["data"]} == {"L Gamma"}

    like_res = await run_report(name="List Over Reports", filters={"report_name__like": "L "})
    assert {r["report_name"] for r in like_res["data"]} == {"L Alpha", "L Beta", "L Gamma"}

    # An empty filter value is a no-op (untouched filter in the UI).
    noop_res = await run_report(name="List Over Reports", filters={"report_type__eq": ""})
    assert len(noop_res["data"]) >= 4


# ── export_report_xlsx ──────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_export_report_xlsx_returns_valid_workbook(ctx):
    """export_xlsx runs the report and returns a real xlsx file."""
    import io

    import openpyxl
    from fastapi.responses import Response

    from grunt.reports.doctypes.Report.report import export_xlsx as export_report_xlsx

    await ctx.new_doc(
        "Report",
        {
            "report_name": "Xlsx Report",
            "report_type": "Query",
            "query": "SELECT 42 as answer",
        },
    )
    await ctx.db._session().commit()

    response = await export_report_xlsx(name="Xlsx Report", filters={})

    assert isinstance(response, Response)
    assert response.media_type == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "Xlsx Report" in response.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(response.body))
    ws = wb.active
    assert ws is not None
    assert ws.cell(row=1, column=1).value == "answer"
    assert ws.cell(row=2, column=1).value == 42
