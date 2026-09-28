"""Report engine — runs Query/Script/List reports."""

from __future__ import annotations

import time
from typing import TYPE_CHECKING, Any, cast

from fastapi import HTTPException
from sqlalchemy import func, select, text

from grunt import _, log

if TYPE_CHECKING:
    from collections.abc import Callable

    from openpyxl.cell import Cell
    from sqlalchemy.ext.asyncio import AsyncSession

    from grunt.auth.doctypes.User.user import User


class ReportEngine:
    async def run(
        self,
        report_name: str,
        filters: dict,
        user: User,
        session: AsyncSession,
    ) -> dict[str, Any]:
        import grunt

        async with grunt.context(session, None, user):
            rows = await grunt.get_list(
                "Report",
                filters={"report_name": report_name},
                fields=[
                    "report_name",
                    "report_type",
                    "query",
                    "script",
                    "columns",
                    "doctype",
                    "conditions",
                    "sort_by",
                    "sort_order",
                    "row_limit",
                ],
                limit=1,
            )
        report = rows[0] if rows else None
        if not report:
            raise HTTPException(
                status_code=404, detail=_("Report “%(name)s” not found") % {"name": report_name}
            )

        report_type = report["report_type"]
        if report_type == "Query":
            return await self._run_query_report({"query": report["query"] or ""}, filters, session)
        if report_type == "Script":
            return await self._run_script_report(
                {"script": report["script"] or "", "columns": report.get("columns")},
                filters,
                user,
                session,
            )
        if report_type == "List":
            doctype = report.get("doctype")
            if not doctype:
                raise HTTPException(400, detail=_("List report requires a DocType"))
            return await self._run_list_report(doctype, report, filters, user, session)
        raise HTTPException(
            status_code=400,
            detail=_("Report type “%(type)s” is not supported") % {"type": report_type},
        )

    # ── Query report ──────────────────────────────────────────────────────────

    async def _run_query_report(
        self,
        report: dict,
        filters: dict,
        session: AsyncSession,
        declared_columns: list[dict] | None = None,
    ) -> dict[str, Any]:
        query_str = report.get("query", "").strip()
        if not query_str:
            raise HTTPException(400, detail=_("No query specified"))

        # Security check — allow only SELECT
        try:
            import sqlparse

            parsed = sqlparse.parse(query_str)
            for stmt in parsed:
                stmt_type = stmt.get_type()
                if stmt_type not in ("SELECT", None):
                    raise HTTPException(400, detail=_("Only SELECT is allowed"))
        except ImportError:
            log.debug("suppressed_expected_error", exc_info=True)

        forbidden = ["drop", "delete", "update", "insert", "create", "alter", "truncate"]
        query_lower = query_str.lower()
        for word in forbidden:
            if word in query_lower:
                raise HTTPException(400, detail=_("Forbidden: %(word)s") % {"word": word.upper()})

        start = time.time()
        result = await session.execute(text(query_str))
        keys = list(result.keys())
        rows = [dict(zip(keys, row, strict=False)) for row in result.fetchall()]
        elapsed = int((time.time() - start) * 1000)

        # A caller (e.g. a Script report that only *picked* the SQL) may declare
        # column metadata — labels, fieldtypes, `total` flag. Match it to the
        # SQL result keys by fieldname; fall back to a plain text column.
        declared = {c["fieldname"]: c for c in (declared_columns or []) if c.get("fieldname")}
        columns = [declared.get(k, {"fieldname": k, "label": k, "fieldtype": "Text"}) for k in keys]
        return {
            "columns": columns,
            "data": rows,
            "meta": {"rows": len(rows), "time_ms": elapsed},
        }

    # ── Script report ─────────────────────────────────────────────────────────

    async def _run_script_report(
        self,
        report: dict,
        filters: dict,
        user: User,
        session: AsyncSession,
    ) -> dict[str, Any]:
        """Execute a Python script in the sandbox and return {columns, data}.

        The sandbox has **no** database or network access. The script must
        satisfy exactly one of two output contracts:

        1. ``result = {"columns": [...], "data": [...]}`` — the script builds
           the rows itself from ``filters`` and plain Python. ``columns`` is a
           list of ``{fieldname, label, fieldtype}``; omit it to derive plain
           text columns from the first row's keys.
        2. ``query = "SELECT ..."`` — the script only *chooses* a read-only
           SQL statement (typically per ``db_dialect``), which the engine
           then runs through the same SELECT-only guard as a Query report.

        Injected globals: ``filters`` (dict) and ``db_dialect`` (str, e.g.
        ``"sqlite"`` / ``"postgresql"``).

        Example — portable "size per table" report::

            if db_dialect == "postgresql":
                query = "SELECT ... pg_total_relation_size(c.oid) ..."
            else:
                query = "SELECT ... FROM dbstat ..."
        """
        import grunt
        from grunt.db.session import get_engine as _engine_factory

        script_src = report.get("script", "").strip()
        if not script_src:
            raise HTTPException(400, detail=_("No script specified"))

        from grunt.scripting.safe_globals import build_safe_globals, compile_script

        # Not plain compile()+exec(): see safe_globals.py's module docstring
        # — a restricted __builtins__ dict alone doesn't stop attribute
        # traversal (e.g. str.__mro__[-1].__subclasses__()) from reaching
        # subprocess.Popen and similar regardless of what's in __builtins__.
        # compile_script() (RestrictedPython) is the actual sandbox; this
        # exec() must only ever run code it produced.
        compiled = compile_script(script_src)
        if compiled.errors:
            raise HTTPException(400, detail="; ".join(compiled.errors))
        assert compiled.code is not None  # no errors → compiled

        engine = await _engine_factory()
        async with grunt.context(session, engine, user):
            extra_globals = build_safe_globals()
            extra_globals["filters"] = filters
            try:
                extra_globals["db_dialect"] = session.bind.dialect.name
            except AttributeError:
                extra_globals["db_dialect"] = ""

            # grunt.result will be set by the script
            _grunt_ns = extra_globals.get("grunt") or extra_globals.get("_grunt")

            start = time.time()
            try:
                exec(compiled.code, extra_globals)
            except HTTPException:
                raise
            except Exception as exc:
                log.error("report.script_error", error=str(exc))
                raise HTTPException(
                    500, detail=_("Script execution error: %(error)s") % {"error": exc}
                ) from exc
            elapsed = int((time.time() - start) * 1000)

            result: Any = extra_globals.get("result") or (
                _grunt_ns.result if _grunt_ns and hasattr(_grunt_ns, "result") else None
            )
            script_query = extra_globals.get("query")

        if not isinstance(result, dict) or "data" not in result:
            # Contract 2: the script only picked a read-only SQL statement.
            if isinstance(script_query, str) and script_query.strip():
                return await self._run_query_report(
                    {"query": script_query},
                    filters,
                    session,
                    declared_columns=report.get("columns"),
                )
            raise HTTPException(
                500,
                detail=(
                    _(
                        "The script must assign `result = {'columns': [...], 'data': [...]}` "
                        "or `query = 'SELECT ...'`"
                    )
                ),
            )

        columns = result.get("columns") or [
            {"fieldname": k, "label": k, "fieldtype": "Text"}
            for k in (result["data"][0].keys() if result["data"] else [])
        ]
        return {
            "columns": columns,
            "data": result["data"],
            "meta": {"rows": len(result["data"]), "time_ms": elapsed},
        }

    # ── List / Dynamic report ─────────────────────────────────────────────────

    async def _run_list_report(
        self,
        doctype: str,
        report: dict,
        filters: dict,
        user: User,
        session: AsyncSession,
    ) -> dict[str, Any]:
        """Run a dynamic aggregate report against a DocType table.

        ``report["columns"]`` is a list of column definitions::

            [
                {"fieldname": "status", "label": "Status", "group_by": true},
                {"fieldname": "amount", "label": "Total Amount",
                 "aggregation": "sum", "fieldtype": "Float"},
                {"fieldname": "id", "label": "Count",
                 "aggregation": "count", "fieldtype": "Int"},
            ]

        Supported aggregations: ``count``, ``sum``, ``avg``, ``min``, ``max``.
        Columns without an aggregation are treated as GROUP BY columns when
        any aggregation column is present; otherwise a plain SELECT is used.
        """
        import grunt
        from grunt.errors import not_found
        from grunt.permissions.rbac import permission_checker

        dt = await grunt.get_meta(doctype)
        if dt is None:
            raise not_found(_("DocType “%(doctype)s” not found") % {"doctype": doctype})
        table = dt.table

        # Permission check
        await permission_checker.require(user, dt, "read")

        col_defs: list[dict] = report.get("columns") or []

        # If no column definitions, fall back to in_list_view fields
        if not col_defs:
            col_defs = [
                {"fieldname": f.fieldname, "label": f.label} for f in dt.get_list_view_fields()
            ]

        agg_map: dict[str, Callable[..., Any]] = {
            "count": func.count,
            "sum": func.sum,
            "avg": func.avg,
            "min": func.min,
            "max": func.max,
        }

        from grunt.reports.list_options import DATE_BUCKETS, compile_conditions, date_bucket

        dialect = session.get_bind().dialect.name
        # "none" (the builder's "no aggregation") is not an aggregation.
        has_aggregation = any((c.get("aggregation") or "").lower() in agg_map for c in col_defs)
        select_cols = []
        group_by_cols = []
        group_fields: list[str] = []
        date_groups: dict[str, str] = {}
        order_exprs: dict[str, Any] = {}  # result column → sortable expression
        result_columns = []

        for col_def in col_defs:
            fn = col_def["fieldname"]
            label = col_def.get("label", fn)
            agg = (col_def.get("aggregation") or "").lower()
            fieldtype = col_def.get("fieldtype", "Text")

            sa_col = table.c.get(fn)
            if sa_col is None:
                continue

            if agg in agg_map:
                expr = agg_map[agg](sa_col)
                select_cols.append(expr.label(fn))
            elif has_aggregation:
                # Non-aggregated column when aggregation is present → GROUP BY,
                # optionally by period for a date column.
                bucket = col_def.get("date_group")
                if bucket in DATE_BUCKETS:
                    expr = date_bucket(sa_col, bucket, dialect)
                    date_groups[fn] = bucket
                    fieldtype = "Data"
                    select_cols.append(expr.label(fn))
                else:
                    expr = sa_col
                    select_cols.append(sa_col)
                group_by_cols.append(expr)
                group_fields.append(fn)
            else:
                expr = sa_col
                select_cols.append(sa_col)
            order_exprs[fn] = expr
            result_columns.append({"fieldname": fn, "label": label, "fieldtype": fieldtype})

        if not select_cols:
            return {"columns": [], "data": [], "meta": {"rows": 0, "time_ms": 0}}

        stmt = select(*select_cols)

        # Apply filters — operator-aware, same vocabulary as list views:
        # ``field__gte`` / ``__like`` / ``__ne`` / … ; a bare ``field`` means
        # equality. Empty values are dropped so an untouched filter is a no-op.
        from grunt.db.filters import build_clauses

        active_filters = {k: v for k, v in filters.items() if v not in (None, "", [])}
        for clause in build_clauses(table, active_filters):
            stmt = stmt.where(clause)

        # Fixed report conditions (always on, unlike the viewer's filters).
        conditions = compile_conditions(report.get("conditions"))
        for condition in conditions:
            for clause in build_clauses(table, condition):
                stmt = stmt.where(clause)

        # Row-level security
        from grunt.permissions.query import apply_permission_filter
        from grunt.permissions.user_permissions import build_conditions

        stmt = await apply_permission_filter(stmt, table, user, dt)
        up_conds = await build_conditions(table, user, dt)
        if up_conds:
            from sqlalchemy import and_

            stmt = stmt.where(and_(*up_conds))

        if group_by_cols:
            stmt = stmt.group_by(*group_by_cols)

        # Sort by any result column; a grouped report defaults to its groups.
        sort_by = report.get("sort_by")
        descending = (report.get("sort_order") or "asc").lower() == "desc"
        if sort_by in order_exprs:
            expr = order_exprs[sort_by]
            stmt = stmt.order_by(expr.desc() if descending else expr.asc())
        elif group_by_cols:
            stmt = stmt.order_by(*group_by_cols)

        # "Top N": the report's own limit, never above the hard cap.
        row_limit = int(report.get("row_limit") or 0)
        stmt = stmt.limit(min(row_limit, 10_000) if row_limit > 0 else 10_000)

        start = time.time()
        db_result = await session.execute(stmt)
        keys = list(db_result.keys())
        rows = [dict(zip(keys, row, strict=False)) for row in db_result.fetchall()]
        elapsed = int((time.time() - start) * 1000)

        meta: dict[str, Any] = {"rows": len(rows), "time_ms": elapsed}
        if group_by_cols:
            # Each result row summarises the documents sharing these group
            # values — the client can open them as a filtered list. Date groups
            # are periods (the client turns the label into a date range), and
            # the report's own conditions narrow that list too.
            meta["drilldown"] = {
                "doctype": doctype,
                "group_by": group_fields,
                "date_groups": date_groups,
                "conditions": {k: v for c in conditions for k, v in c.items()},
            }

        return {"columns": result_columns, "data": rows, "meta": meta}

    # ── Excel export ──────────────────────────────────────────────────────────

    async def export_excel(self, result: dict, report_name: str) -> bytes:
        try:
            from io import BytesIO

            import openpyxl
            from openpyxl.styles import Font
        except ImportError as e:
            raise HTTPException(status_code=500, detail=_("openpyxl is not installed")) from e

        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = report_name[:31]

        columns = result.get("columns", [])
        data = result.get("data", [])

        # Header row
        for ci, col in enumerate(columns, 1):
            cell = ws.cell(row=1, column=ci, value=col.get("label", col["fieldname"]))
            cell.font = Font(bold=True)

        # Data rows
        from grunt.io.exporters.sanitize import escape_formula

        for ri, row in enumerate(data, 2):
            for ci, col in enumerate(columns, 1):
                cell_value = row.get(col["fieldname"])
                if isinstance(cell_value, str):
                    cell_value = escape_formula(cell_value)
                ws.cell(row=ri, column=ci, value=cell_value)

        # Totals row — sum every column flagged with `total: true`.
        total_fields = {c["fieldname"] for c in columns if c.get("total")}
        if total_fields and data:
            tr = len(data) + 2
            for ci, col in enumerate(columns, 1):
                fn = col["fieldname"]
                if fn in total_fields:
                    value: Any = sum(row.get(fn) or 0 for row in data)
                elif ci == 1:
                    value = _("Total")
                else:
                    continue
                cell = ws.cell(row=tr, column=ci, value=value)
                cell.font = Font(bold=True)

        # Auto-width
        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=0)
            ws.column_dimensions[cast("Cell", col[0]).column_letter].width = min(max_len + 2, 50)

        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()


report_engine = ReportEngine()
