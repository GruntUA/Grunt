"""Regression: CSV/Excel formula injection (CWE-1236) in document exports.

A document field value under attacker control (any regular text field) that
starts with =, +, -, @ (or a leading tab/CR) is auto-interpreted as a live
formula by Excel/LibreOffice/Sheets when the exported file is opened —
verified live against openpyxl: writing the string
'=HYPERLINK("http://evil.com","click")' as a cell value produces
cell.data_type == "f" (formula), not plain text. That's real formula
execution — e.g. data exfiltration via HYPERLINK() — in whoever (often an
admin) opens an export of user-submitted data.
"""

import io

import openpyxl
import pytest

from grunt.io.exporters.csv import CsvExporter
from grunt.io.exporters.sanitize import escape_formula
from grunt.io.exporters.xlsx import XlsxExporter
from grunt.metadata.field import DocField


def _field(fieldname: str, label: str) -> DocField:
    return DocField(fieldname=fieldname, label=label, fieldtype="Data")


DANGEROUS_VALUES = [
    '=HYPERLINK("http://evil.example.com","click me")',
    "+cmd|'/c calc'!A0",
    "-2+3",
    "@SUM(1,1)",
    "\t=1+1",
]


class TestEscapeFormula:
    @pytest.mark.parametrize("value", DANGEROUS_VALUES)
    def test_dangerous_values_get_quote_prefixed(self, value):
        result = escape_formula(value)
        assert result.startswith("'")
        assert result == f"'{value}"

    @pytest.mark.parametrize("value", ["hello", "user@example.com", "100", ""])
    def test_safe_values_unchanged(self, value):
        assert escape_formula(value) == value


class TestCsvExporterEscapesFormulas:
    @pytest.mark.asyncio
    async def test_formula_value_escaped_in_csv_output(self):
        exporter = CsvExporter()
        fields = [_field("notes", "Notes")]
        rows = [{"notes": '=HYPERLINK("http://evil.example.com","click")'}]

        content = await exporter.export("TestDoc", rows, fields)
        text = content.decode("utf-8-sig")

        assert "'=HYPERLINK" in text
        # The raw, unescaped formula must never appear as a bare leading value.
        assert '\n=HYPERLINK("http://evil.example.com","click")' not in text


class TestXlsxExporterEscapesFormulas:
    @pytest.mark.asyncio
    async def test_formula_value_is_not_a_live_formula_cell(self):
        exporter = XlsxExporter()
        fields = [_field("notes", "Notes")]
        rows = [{"notes": '=HYPERLINK("http://evil.example.com","click")'}]

        content = await exporter.export("TestDoc", rows, fields)

        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
        assert ws is not None
        cell = ws.cell(row=2, column=1)
        assert cell.data_type != "f", "value must not be interpreted as a live formula"
        assert isinstance(cell.value, str) and cell.value.startswith("'")

    @pytest.mark.asyncio
    async def test_numbers_and_dates_pass_through_untouched(self):
        """Only strings need escaping — numeric/date cells are never
        interpreted as formulas by openpyxl regardless of value.
        """
        exporter = XlsxExporter()
        fields = [_field("amount", "Amount")]
        rows = [{"amount": 42}]

        content = await exporter.export("TestDoc", rows, fields)
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
        assert ws is not None
        cell = ws.cell(row=2, column=1)
        assert cell.value == 42
